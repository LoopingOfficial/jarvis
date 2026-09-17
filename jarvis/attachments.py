"""Pièces jointes JARVIS — stockage temporaire sécurisé + extraction réelle.

Chaque fichier envoyé depuis la Command Bar est :

1. validé (extension en liste blanche, taille max, signature binaire réelle) ;
2. stocké sous un nom UUID dans ``data/uploads`` — jamais le nom utilisateur ;
3. jamais exécuté (lecture/analyse uniquement) ;
4. nettoyé après un TTL (24 h par défaut).

``extract()`` produit un contenu réellement lisible par le LLM selon le type
réel du fichier (texte/code, CSV, PDF, XLSX, DOCX, image).
"""
from __future__ import annotations

import base64
import csv
import datetime as _dt
import io
import json
import mimetypes
import re
import time
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from .config import DATA_DIR

UPLOAD_DIR = Path(DATA_DIR) / "uploads"

# ---------------------------------------------------------------------------
# Formats réellement analysables.
# ---------------------------------------------------------------------------
IMAGE_EXTS = {".jpg", ".jpeg", ".png", ".webp"}
PDF_EXTS = {".pdf"}
TEXT_EXTS = {
    ".txt", ".md", ".json", ".xml", ".log", ".csv",
    # Transcripts d'appels et sous-titres (document.parse_transcript).
    ".vtt", ".srt",
    # code
    ".js", ".mjs", ".ts", ".jsx", ".tsx", ".php", ".py", ".html", ".css",
    ".scss", ".sql", ".yaml", ".yml", ".toml", ".ini", ".cfg",
    ".sh", ".bash", ".ps1", ".bat", ".cmd",
    ".java", ".c", ".cpp", ".h", ".hpp", ".rb", ".go", ".rs", ".swift", ".kt",
    ".cs",
}
XLSX_EXTS = {".xlsx"}
DOCX_EXTS = {".docx"}

ALLOWED_EXTS = IMAGE_EXTS | PDF_EXTS | TEXT_EXTS | XLSX_EXTS | DOCX_EXTS

# Exécutables binaires / archives : refusés quel que soit le nom.
HARD_DENY_EXTS = {
    ".exe", ".dll", ".msi", ".com", ".scr", ".sys", ".drv",
    ".zip", ".rar", ".7z", ".tar", ".gz", ".bz2", ".xz", ".iso", ".img",
    ".bin", ".dat", ".whl", ".apk", ".ipa", ".msix", ".jar", ".class",
    ".so", ".dylib", ".pyc", ".pyd",
}

KIND_LABELS = {"image": "Image", "pdf": "PDF", "text": "Texte/Code",
               "csv": "CSV", "xlsx": "XLSX", "docx": "DOCX"}

_TRAVERSAL = re.compile(r"[\\/\x00-\x1f]")

_IMAGE_MAGIC = {
    b"\x89PNG\r\n\x1a\n": "png",
    b"\xff\xd8\xff": "jpg",
    b"RIFF": "webp",   # confirmé plus bas par "WEBP"
}


class AttachmentError(Exception):
    """Erreur lisible par l'utilisateur (jamais une stacktrace)."""


@dataclass
class Attachment:
    id: str
    filename: str
    stored_name: str
    ext: str
    mime: str
    kind: str
    size: int
    created_at: float
    path: Path
    width: int = 0
    height: int = 0

    def public(self) -> dict[str, Any]:
        out = {
            "id": self.id, "filename": self.filename, "ext": self.ext,
            "mime": self.mime, "kind": self.kind, "size": self.size,
            "created_at": self.created_at, "kind_label": KIND_LABELS.get(self.kind, self.kind),
            "url": f"/api/uploads/{self.id}/file",
            "thumb": f"/api/uploads/{self.id}/thumb" if self.kind == "image" else "",
        }
        if self.kind == "image" and self.width:
            out["width"], out["height"] = self.width, self.height
        return out


class AttachmentStore:
    def __init__(self, core) -> None:
        self.core = core
        self._dir = UPLOAD_DIR
        self._dir.mkdir(parents=True, exist_ok=True)
        self._index: dict[str, Attachment] = {}
        self._extract_cache: dict[str, dict[str, Any]] = {}
        self._scan()

    # -- réglages ----------------------------------------------------------
    def _max_size(self, kind: str) -> int:
        section = "files"
        mb = int(self.core.settings.get(section, "max_upload_size_mb", 25))
        if kind == "image":
            mb = int(self.core.settings.get(section, "max_image_size_mb", mb))
        return mb * 1024 * 1024

    def _max_text_chars(self) -> int:
        return int(self.core.settings.get("files", "text_context_chars", 30_000))

    def _max_pdf_pages(self) -> int:
        return int(self.core.settings.get("files", "pdf_page_context", 25))

    # -- index -------------------------------------------------------------
    def _scan(self) -> None:
        self._index.clear()
        try:
            entries = list(self._dir.iterdir())
        except OSError:
            return
        for sidecar in entries:
            if sidecar.suffix != ".json" or sidecar.name == "index.json":
                continue
            try:
                data = json.loads(sidecar.read_text(encoding="utf-8"))
                if not isinstance(data, dict):
                    continue
                aid = str(data.get("id") or "")
                stored = str(data.get("stored_name") or "")
                path = self._dir / stored
                if not aid or not path.is_file():
                    continue
                att = Attachment(
                    id=aid, filename=str(data.get("filename") or stored),
                    stored_name=stored, ext=str(data.get("ext") or ""),
                    mime=str(data.get("mime") or ""), kind=str(data.get("kind") or "text"),
                    size=int(data.get("size") or 0), created_at=float(data.get("created_at") or 0),
                    path=path, width=int(data.get("width") or 0), height=int(data.get("height") or 0),
                )
                self._index[aid] = att
            except Exception:
                continue

    def _save_meta(self, att: Attachment) -> None:
        side = self._dir / f"{att.id}.json"
        side.write_text(json.dumps({
            "id": att.id, "filename": att.filename, "stored_name": att.stored_name,
            "ext": att.ext, "mime": att.mime, "kind": att.kind, "size": att.size,
            "created_at": att.created_at, "width": att.width, "height": att.height,
        }, ensure_ascii=False), encoding="utf-8")

    def list(self, limit: int = 50) -> list[dict[str, Any]]:
        items = sorted(self._index.values(), key=lambda a: a.created_at, reverse=True)
        return [a.public() for a in items[:limit]]

    def get(self, attachment_id: str) -> Attachment | None:
        return self._index.get(attachment_id)

    def require(self, attachment_id: str) -> Attachment:
        att = self._index.get(attachment_id)
        if att is None or not att.path.is_file():
            raise AttachmentError("Cette pièce jointe n'est plus disponible (supprimée ou expirée).")
        return att

    # -- ingestion sécurisée ----------------------------------------------
    @staticmethod
    def _safe_extension(filename: str) -> str:
        """Extension en liste blanche, sinon erreur humainement lisible."""
        name = str(filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
        ext = Path(name or "fichier").suffix.lower() or ""
        if ext in HARD_DENY_EXTS or ext not in ALLOWED_EXTS:
            label = ext if ext not in HARD_DENY_EXTS and ext else "(inconnue)"
            raise AttachmentError(
                f"Format non supporté ({label or 'sans extension'}). J'envoie mes analyses "
                f"sur : images (jpg, png, webp), texte/code (txt, md, js, php, py, html, "
                f"css, sql, csv, log, json, xml, yaml…), PDF, XLSX et DOCX.")
        return ext

    @staticmethod
    def _sanitize_display_name(filename: str) -> str:
        name = str(filename or "").replace("\\", "/").rsplit("/", 1)[-1].strip()
        name = _TRAVERSAL.sub("", name)[:180]
        return name or "fichier"

    def _detect_kind(self, ext: str, raw: bytes) -> str:
        if ext in IMAGE_EXTS: return "image"
        if ext in PDF_EXTS: return "pdf"
        if ext in XLSX_EXTS: return "xlsx"
        if ext in DOCX_EXTS: return "docx"
        if ext == ".csv": return "csv"
        return "text"

    def _verify_signature(self, att: Attachment, raw: bytes) -> None:
        """Vérifie le type RÉEL, pas la seule extension."""
        if att.kind == "image":
            verified = self._image_signature(raw)
            if verified is None:
                raise AttachmentError("Le fichier ne ressemble pas à une image valide (corrompu ou renommé).")
            detected = IMAGE_EXTS | {".jpg", ".jpeg"}
            if verified not in {".png", ".jpg", ".jpeg", ".webp"}:
                raise AttachmentError("Format d'image non pris en charge : " + (verified or att.ext))
        elif att.kind == "pdf":
            if raw[:5] != b"%PDF-":
                raise AttachmentError("Je n'arrive pas à lire ce fichier PDF (en-tête invalide).")
        elif att.kind in {"xlsx", "docx"}:
            if raw[:2] != b"PK":
                raise AttachmentError("Le document est corrompu ou n'est pas un fichier valide.")
        elif att.kind == "text":
            if b"\x00" in raw[:8192] or not self._try_decode(raw[:4096]):
                raise AttachmentError("Ce fichier texte semble binaire ; je ne peux pas l'analyser comme du texte.")

    @staticmethod
    def _image_signature(raw: bytes) -> str | None:
        if raw[:8] == b"\x89PNG\r\n\x1a\n":
            return ".png"
        if raw[:3] == b"\xff\xd8\xff":
            return ".jpg"
        if raw[:4] == b"RIFF" and raw[8:12] == b"WEBP":
            return ".webp"
        return None

    def _image_dimensions(self, raw: bytes) -> tuple[int, int]:
        try:
            from PIL import Image
            with Image.open(io.BytesIO(raw)) as img:
                return img.size
        except Exception:
            return 0, 0

    def store(self, filename: str, data: bytes, *, mime_hint: str = "") -> Attachment:
        """Valide et enregistre un upload. Lève ``AttachmentError`` lisible."""
        data = data or b""
        ext = self._safe_extension(filename)
        display = self._sanitize_display_name(filename)
        kind = self._detect_kind(ext, data)
        if not data:
            raise AttachmentError("Le fichier est vide.")
        if len(data) > self._max_size(kind):
            raise AttachmentError(
                f"Fichier trop volumineux (max {self._max_size(kind) // 1024 // 1024} Mo "
                f"pour {KIND_LABELS.get(kind, kind)}).")

        aid = "att_" + uuid.uuid4().hex[:12]
        stored = f"{uuid.uuid4().hex}{ext}"
        safe_mime = mimetypes.guess_type(display)[0] or mime_hint or "application/octet-stream"
        att = Attachment(id=aid, filename=display, stored_name=stored, ext=ext,
                         mime=safe_mime, kind=kind, size=len(data),
                         created_at=time.time(), path=self._dir / stored)
        self._verify_signature(att, data)
        # L'extension stockée doit refléter le type réel détecté (anti-faux nom).
        if att.kind == "image" and self._image_signature(data):
            real = self._image_signature(data)
            if not att.path.suffix.lower() == real:
                stored = f"{uuid.uuid4().hex}{real}"
                att.path = self._dir / stored
                att.ext = real
        att.path.write_bytes(data)
        if att.kind == "image":
            att.width, att.height = self._image_dimensions(data)
        self._save_meta(att)
        self._index[aid] = att
        return att

    def ingest_file(self, path: str | Path, filename: str = "") -> Attachment:
        """Importe un fichier déjà présent sur la machine (ex. téléchargement navigateur)."""
        src = Path(path)
        raw = src.read_bytes()
        return self.store(filename or src.name, raw)

    # -- suppression / nettoyage ------------------------------------------
    def delete(self, attachment_id: str) -> bool:
        att = self._index.pop(attachment_id, None)
        if att is None:
            return False
        for p in (att.path, self._dir / f"{att.id}.json", self._dir / f"{att.id}_thumb.png"):
            try:
                p.unlink(missing_ok=True)
            except OSError:
                pass
        self._extract_cache.pop(attachment_id, None)
        return True

    def cleanup(self) -> int:
        """Supprime les pièces jointes plus vieilles que le TTL. Retourne le compte."""
        ttl_h = int(self.core.settings.get("files", "attachment_ttl_hours", 24))
        cutoff = time.time() - ttl_h * 3600
        removed = 0
        for aid in list(self._index):
            att = self._index[aid]
            if att.created_at < cutoff:
                self.delete(aid)
                removed += 1
        return removed

    # -- décodage texte ----------------------------------------------------
    @staticmethod
    def _try_decode(data: bytes) -> str | None:
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                return data.decode(enc)
            except (UnicodeDecodeError, ValueError):
                continue
        return None

    def _read_text(self, att: Attachment) -> str:
        raw = att.path.read_bytes()
        text = self._try_decode(raw) or ""
        return text

    # -- extraction ---------------------------------------------------------
    def extract(self, att: Attachment) -> dict[str, Any]:
        if att.id in self._extract_cache:
            return self._extract_cache[att.id]
        try:
            out = self._extract_uncached(att)
        except AttachmentError:
            raise
        except Exception as exc:
            raise AttachmentError(f"Impossible de lire ce fichier ({type(exc).__name__}).") from exc
        out["_when"] = time.time()
        self._extract_cache[att.id] = out
        return out

    def _extract_uncached(self, att: Attachment) -> dict[str, Any]:
        if att.kind == "image":
            return {"kind": "image", "width": att.width, "height": att.height}
        if att.kind == "pdf":
            return self._extract_pdf(att)
        if att.kind == "csv":
            return self._extract_csv(att)
        if att.kind == "xlsx":
            return self._extract_xlsx(att)
        if att.kind == "docx":
            return self._extract_docx(att)
        return self._extract_code(att)

    def _extract_code(self, att: Attachment) -> dict[str, Any]:
        text = self._read_text(att)
        limit = self._max_text_chars()
        head = text[:limit]
        return {"kind": "text", "text": head, "chars": len(text),
                "truncated": len(text) > len(head)}

    def _extract_pdf(self, att: Attachment) -> dict[str, Any]:
        from pypdf import PdfReader
        reader = PdfReader(att.path)
        page_count = len(reader.pages)
        pages: list[dict[str, Any]] = []
        for i in range(min(page_count, self._max_pdf_pages())):
            try:
                text = (reader.pages[i].extract_text() or "").strip()
            except Exception:
                text = ""
            pages.append({"page": i + 1, "text": text[:20_000]})
        joined = "\n\n".join(f"[Page {p['page']}]\n{p['text']}" for p in pages)
        total_chars = sum(len(p["text"]) for p in pages)
        text_only = "".join(p["text"] for p in pages).strip()
        return {
            "kind": "pdf", "page_count": page_count,
            "pages": pages, "text": joined, "chars": total_chars,
            "truncated": page_count > len(pages),
            "text_extracted": bool(text_only),
            "scanned_likely": not text_only,
        }

    def _extract_csv(self, att: Attachment) -> dict[str, Any]:
        raw = att.path.read_bytes()
        text = self._try_decode(raw) or ""
        text = text.lstrip("\ufeff")
        dialect_delim = self._sniff_delimiter(text)
        sample_rows, row_count, col_types, header = self._parse_csv_rows(text, dialect_delim)
        return {
            "kind": "csv", "delimiter": dialect_delim,
            "header": header, "columns": header,
            "row_count": row_count, "sample": sample_rows,
            "col_types": col_types,
        }

    @staticmethod
    def _sniff_delimiter(text: str) -> str:
        first = text.splitlines()[0] if text.splitlines() else ""
        counters = {",": 0, ";": 0, "\t": 0, "|": 0}
        for line in text.splitlines()[:20]:
            for d in counters:
                counters[d] += line.count(d)
        best = max(counters, key=counters.get)
        return best if counters[best] > 0 else ","

    def _parse_csv_rows(self, text: str, delimiter: str):
        sample_rows: list[list[str]] = []
        row_count = 0
        values_by_col: list[list[str]] = []
        header: list[str] = []
        try:
            reader = csv.reader(io.StringIO(text), delimiter=delimiter)
            for ri, row in enumerate(reader):
                row = [c.strip().replace("\n", " ") for c in row]
                if ri == 0:
                    header = row or []
                    if ri < 16:
                        sample_rows.append(row)
                    row_count = 1
                    continue
                for ci, value in enumerate(row):
                    while len(values_by_col) <= ci:
                        values_by_col.append([])
                    if value:
                        values_by_col[ci].append(value)
                if ri < 16:
                    sample_rows.append(row)
                row_count += 1
        except Exception:
            pass
        col_types: dict[str, str] = {}
        for ci, values in enumerate(values_by_col):
            name = header[ci] if ci < len(header) else f"colonne {ci + 1}"
            col_types[str(name)] = AttachmentStore._infer_csv_type(values[:400])
        return sample_rows[1:] if sample_rows else [], row_count, col_types, header

    @staticmethod
    def _infer_csv_type(values: list[str]) -> str:
        if not values:
            return "vide"
        nonempty = [v for v in values if v.strip()]
        if not nonempty:
            return "vide"
        if all(AttachmentStore._is_int(v) for v in nonempty):
            return "entier"
        if all(AttachmentStore._is_float(v) for v in nonempty):
            return "nombre"
        if all(AttachmentStore._is_date(v) for v in nonempty):
            return "date"
        return "texte"

    @staticmethod
    def _is_int(v: str) -> bool:
        return bool(re.fullmatch(r"[-+]?\d{1,12}", v.strip()))

    @staticmethod
    def _is_float(v: str) -> bool:
        return bool(re.fullmatch(r"[-+]?\d+(?:[.,]\d+)?(?:[eE][-+]?\d+)?", v.strip().replace("\u202f", "")))

    @staticmethod
    def _is_date(v: str) -> bool:
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%Y/%m/%d"):
            try:
                _dt.datetime.strptime(v.strip(), fmt)
                return True
            except ValueError:
                continue
        return False

    def _extract_xlsx(self, att: Attachment) -> dict[str, Any]:
        from openpyxl import load_workbook
        wb = load_workbook(att.path, read_only=True, data_only=True)
        sheets: list[dict[str, Any]] = []
        for ws in wb.worksheets:
            rows_iter = ws.iter_rows(values_only=True)
            sample: list[list[str]] = []
            row_count = 0
            max_cols = 0
            for ri, row in enumerate(rows_iter):
                if ri == 0:
                    raw = ["" if v is None else str(v) for v in row]
                    header = raw
                values = ["" if v is None else str(v) for v in row]
                if ri < 15:
                    sample.append(values)
                row_count += 1
                max_cols = max(max_cols, len(values))
            sheets.append({
                "name": ws.title, "rows": row_count, "cols": max_cols,
                "header": header if row_count else [],
                "sample": sample[:14] if row_count else [],
            })
        wb.close()
        return {"kind": "xlsx", "sheet_count": len(sheets) or 0, "sheets": sheets}

    def _extract_docx(self, att: Attachment) -> dict[str, Any]:
        from docx import Document
        doc = Document(att.path)
        paragraphs = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        table_count = len(doc.tables)
        if table_count:
            for table in doc.tables[:3]:
                for row in list(table.rows)[:8]:
                    paragraphs.append(" | ".join(cell.text.strip().replace("\n", " ") for cell in row.cells))
        text = "\n".join(paragraphs)
        limit = self._max_text_chars()
        return {"kind": "docx", "paragraph_count": len(paragraphs),
                "table_count": table_count, "text": text[:limit],
                "truncated": len(text) > limit}

    # -- contexte LLM -------------------------------------------------------
    def context_block(self, att: Attachment, question: str = "") -> dict[str, Any]:
        """Bloc de contexte prêt pour le LLM.

        Retourne ``{"kind", "text", "vision": bool, "attachment": att}``.
        ``vision=True`` : l'image doit être analysée par un modèle VISION
        (jamais ComfyUI) et le résultat intégré à la même requête.
        """
        if att.kind == "image":
            raw = base64.b64encode(att.path.read_bytes()).decode("ascii")
            return {"kind": "image", "text": "", "vision": True,
                    "attachment": att, "base64": raw}
        x = self.extract(att)
        header = (f"\n<FICHIER JOINT> {att.filename} ({KIND_LABELS.get(att.kind, att.kind)}, "
                  f"{att.size // 1024} Ko)\n")
        if att.kind == "text":
            body = x["text"] or "(fichier vide)"
            if x.get("truncated"):
                body += (f"\n… (contenu tronqué : {x['chars']} caractères au total, "
                         f"seuls les {self._max_text_chars()} premiers sont injectés)")
            return {"kind": "text", "text": header + "CONTENU\n" + body, "vision": False}
        if att.kind == "csv":
            cols = ", ".join(x["columns"]) if x["columns"] else "(aucune colonne détectée)"
            lines = [
                header + f"STRUCTURE CSV : {x['row_count']} ligne(s) de données",
                f"délimiteur = {x['delimiter']}",
                f"colonnes = {cols}",
                "types par colonne : " + ", ".join(f"{k} → {v}" for k, v in x["col_types"].items()),
                "EXTRAIT (5 premières lignes) :",
            ]
            for row in x["sample"][:5]:
                lines.append("- " + " | ".join(str(c)[:60] for c in row))
            lines.append("Jamais d'invention : les cellules non présentes ci-dessus restent inconnues.")
            return {"kind": "text", "text": "\n".join(lines), "vision": False}
        if att.kind == "pdf":
            if not x.get("text_extracted"):
                return {"kind": "text", "text": (
                    header + "Aucun texte extractible : ce PDF est probablement constitué "
                    "d'images/scans. Le rendu OCR n'étant pas installé, je ne peux pas en lire "
                    "le contenu — mais je peux t'indiquer comment le convertir."),
                    "vision": False}
            body = x["text"]
            if x.get("truncated"):
                body += (f"\n… (PDF de {x['page_count']} pages, seules les "
                         f"{self._max_pdf_pages()} premières pages sont injectées)")
            return {"kind": "text", "text": header + f"PAGES ({x['page_count']})\n" + body,
                    "vision": False}
        if att.kind == "xlsx":
            lines = [header + f"CLASSEUR : {x['sheet_count']} feuille(s). Onglets RÉELS :"]
            for s in x["sheets"]:
                lines.append(f"- « {s['name']} » : {s['rows']} ligne(s), {s['cols']} colonne(s)")
                if s["header"]:
                    lines.append("     colonnes : " + ", ".join(s["header"]))
                for row in s["sample"][:4]:
                    lines.append("     - " + " | ".join(str(c)[:40] for c in row))
            lines.append("Interdiction : ne jamais inventer le nom d'un onglet. "
                         "Pour analyser une feuille précise, l'utilisateur doit la nommer.")
            return {"kind": "text", "text": "\n".join(lines), "vision": False}
        if att.kind == "docx":
            body = x["text"] or "(document sans texte)"
            if x.get("truncated"):
                body += f"\n… (contenu tronqué, {x['paragraph_count']} paragraphes)"
            return {"kind": "text", "text": header + "CONTENU\n" + body, "vision": False}
        return {"kind": "text", "text": header + "(contenu non extractible)", "vision": False}

    # -- service fichier / miniature ----------------------------------------
    def serve(self, attachment_id: str) -> tuple[bytes, str] | None:
        att = self.require(attachment_id)
        data = att.path.read_bytes()
        if att.kind == "image":
            ctype = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg",
                     "webp": "image/webp"}.get(att.ext, "application/octet-stream")
        elif att.kind == "pdf":
            ctype = "application/pdf"
        elif att.ext == ".csv":
            ctype = "text/csv"
        else:
            # HTML/SVG et code : servis en texte brut, jamais actifs.
            ctype = "text/plain; charset=utf-8"
        return data, ctype

    def thumbnail(self, attachment_id: str, box: int = 240) -> tuple[bytes, str] | None:
        att = self.require(attachment_id)
        if att.kind != "image":
            return None
        cached = self._dir / f"{att.id}_thumb.png"
        if cached.is_file() and cached.stat().st_mtime >= att.created_at:
            return cached.read_bytes(), "image/png"
        try:
            from PIL import Image
            with Image.open(att.path) as img:
                img.thumbnail((box, box))
                buf = io.BytesIO()
                img.convert("RGB").save(buf, format="PNG")
        except Exception:
            return None
        try:
            cached.write_bytes(buf.getvalue())
        except OSError:
            pass
        return buf.getvalue(), "image/png"

    def resolve(self, ids: list[str]) -> dict[str, Any]:
        """Résout une liste d'identifiants : ``{attachments, errors}``."""
        attachments: list[dict[str, Any]] = []
        errors: list[str] = []
        for raw in (ids or []):
            att = self.get(str(raw))
            if att is None:
                errors.append(f"Pièce jointe « {raw} » introuvable (expirée ?).")
            else:
                attachments.append({"attachment": att.public(),
                                    "context": self.context_block(att)})
        return {"attachments": attachments, "errors": errors}

    def capabilities(self) -> dict[str, Any]:
        vision = self.core.llm.vision_status() if hasattr(self.core, "llm") else {}
        return {
            "vision": vision,
            "formats": {
                "images": sorted(e.lstrip(".") for e in IMAGE_EXTS),
                "text": sorted(e.lstrip(".") for e in TEXT_EXTS),
                "documents": ["pdf", "xlsx", "docx", "csv"],
            },
            "max_upload_bytes": self._max_size("text"),
            "max_attachments": int(self.core.settings.get("files", "max_attachments_per_message", 6)),
        }